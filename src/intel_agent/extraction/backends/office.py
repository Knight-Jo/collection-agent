"""Office structure backend: DOCX, PPTX, XLSX (spec §8.5)."""

from __future__ import annotations

import io
import zipfile

from ...contracts.documents import CoverageUnit, Locator
from ..models import Availability, BackendRequest, BackendOutput
from ._base import BaseBackend
from ._util import make_block

OFFICE_MEDIA = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class OfficeBackend(BaseBackend):
    backend_id = "office"
    version = "1"
    capabilities = ("docx_structure", "pptx_structure", "xlsx_structure")
    media_types = tuple(sorted(OFFICE_MEDIA))

    def availability(self) -> Availability:
        try:
            import docx  # noqa: F401
            import openpyxl  # noqa: F401
            import pptx  # noqa: F401
        except ImportError:
            return "missing_dependency"
        return "available"

    async def run(self, request: BackendRequest) -> BackendOutput:
        data = self.read_bytes(request.resource_id)
        if request.capability == "docx_structure":
            return await self._docx(data)
        if request.capability == "pptx_structure":
            return await self._pptx(data)
        if request.capability == "xlsx_structure":
            return await self._xlsx(data)
        return BackendOutput(warnings=["unsupported capability"])

    async def _docx(self, data: bytes) -> BackendOutput:
        import docx
        from docx.document import Document as _Doc
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(io.BytesIO(data))
        blocks = []
        ordinal = 1
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                para = Paragraph(child, document)
                text = para.text.strip()
                if text:
                    style = para.style.name if para.style else ""
                    block_type = "heading" if style.startswith(
                        "Heading"
                    ) else "paragraph"
                    blocks.append(
                        make_block(
                            self.backend_id, self.version, ordinal, text,
                            block_type,
                            locator=Locator(
                                section_path=[style] if style else []
                            ),
                        )
                    )
                    ordinal += 1
            elif child.tag.endswith("}tbl"):
                table = Table(child, document)
                rows = [
                    " | ".join(c.text.strip() for c in row.cells)
                    for row in table.rows
                ]
                blocks.append(
                    make_block(
                        self.backend_id, self.version, ordinal,
                        "\n".join(rows), "table",
                    )
                )
                ordinal += 1
        return BackendOutput(
            blocks=blocks,
            coverage=[
                CoverageUnit(
                    unit_type="document", locator=Locator(),
                    status="success" if blocks else "empty",
                )
            ],
        )

    async def _pptx(self, data: bytes) -> BackendOutput:
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(data))
        blocks = []
        ordinal = 1
        for slide_num, slide in enumerate(presentation.slides, start=1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if text:
                    blocks.append(
                        make_block(
                            self.backend_id, self.version, ordinal, text,
                            "paragraph",
                            locator=Locator(slide=slide_num),
                        )
                    )
                    ordinal += 1
        return BackendOutput(
            blocks=blocks,
            coverage=[
                CoverageUnit(
                    unit_type="document", locator=Locator(),
                    status="success" if blocks else "empty",
                )
            ],
        )

    async def _xlsx(self, data: bytes) -> BackendOutput:
        import openpyxl

        cached = openpyxl.load_workbook(
            io.BytesIO(data), data_only=True, read_only=True
        )
        formulas = openpyxl.load_workbook(
            io.BytesIO(data), data_only=False, read_only=True
        )
        blocks = []
        warnings: list[str] = []
        ordinal = 1
        for sheet_name in cached.sheetnames:
            cached_sheet = cached[sheet_name]
            formula_sheet = formulas[sheet_name]
            for row in cached_sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    formula_cell = formula_sheet[cell.coordinate]
                    formula = (
                        formula_cell.value
                        if isinstance(formula_cell.value, str)
                        and formula_cell.value.startswith("=")
                        else None
                    )
                    cached_value = value
                    if formula is not None and cached_value is None:
                        warnings.append("uncalculated_formula")
                    blocks.append(
                        make_block(
                            self.backend_id, self.version, ordinal,
                            str(value), "cell",
                            locator=Locator(
                                sheet=sheet_name,
                                cell_range=cell.coordinate,
                            ),
                            metadata={
                                "formula": formula,
                                "cached_value": value,
                                "sheet": sheet_name,
                            },
                        )
                    )
                    ordinal += 1
        return BackendOutput(
            blocks=blocks,
            coverage=[
                CoverageUnit(
                    unit_type="document", locator=Locator(),
                    status="success" if blocks else "empty",
                )
            ],
            warnings=warnings,
        )
