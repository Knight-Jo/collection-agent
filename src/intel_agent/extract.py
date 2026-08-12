"""Text and link extraction for resources archived by the crawler."""

from __future__ import annotations

import multiprocessing
import os
import queue
import re
import shutil
import signal
import subprocess
import tempfile
import threading
import time
import zipfile
from html.parser import HTMLParser
from importlib import import_module
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

from .document_extract import decode_body, extract_docx_text, extract_html
from .models import ExtractionResult

_REJECTED_SUFFIXES = {
    ".7z",
    ".apk",
    ".bat",
    ".bz2",
    ".cmd",
    ".com",
    ".dll",
    ".dmg",
    ".exe",
    ".gz",
    ".iso",
    ".jar",
    ".js",
    ".msi",
    ".ps1",
    ".rar",
    ".sh",
    ".svg",
    ".tar",
    ".xz",
    ".zip",
}
_REJECTED_MIMES = {
    "application/gzip",
    "application/javascript",
    "application/x-7z-compressed",
    "application/x-bzip2",
    "application/x-dosexec",
    "application/x-msdownload",
    "application/x-rar-compressed",
    "application/x-sh",
    "application/x-tar",
    "application/zip",
    "image/svg+xml",
    "text/javascript",
}
_PLAIN_SUFFIXES = {".csv", ".txt"}
_IMAGE_SUFFIXES = {".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"}
_AUDIO_SUFFIXES = {".flac", ".m4a", ".mp3", ".ogg", ".wav"}
_VIDEO_SUFFIXES = {".mov", ".mp4", ".webm"}
_IMAGE_MIMES = {"image/jpeg", "image/png", "image/tiff", "image/webp"}
_AUDIO_MIMES = {
    "audio/flac",
    "audio/mp4",
    "audio/mpeg",
    "audio/ogg",
    "audio/wav",
    "audio/x-flac",
    "audio/x-m4a",
    "audio/x-wav",
}
_VIDEO_MIMES = {"video/mp4", "video/quicktime", "video/webm"}
_GENERIC_MIMES = {"", "application/octet-stream", "binary/octet-stream"}
_OFFICE_MIMES = {
    "application/msword": "doc",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.ms-excel": "xls",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-powerpoint": "ppt",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
}
_MAX_OOXML_MEMBERS = 1_024
_MAX_OOXML_EXPANDED_BYTES = 32 * 1024 * 1024
_MAX_PDF_PAGES = 200
_MAX_PDF_PAGE_PIXELS = 25_000_000
_MAX_PDF_RASTER_PIXELS = 100_000_000
_MAX_SPREADSHEET_SHEETS = 100
_MAX_SPREADSHEET_CELLS = 250_000
_MAX_PRESENTATION_SLIDES = 500
_MAX_IMAGE_PIXELS = 40_000_000
_EXTRACTION_TIMEOUT_SECONDS = 60
_RESOURCE_WORKER_ACTIVE = False


class _LinkParser(HTMLParser):
    def __init__(self, base_url: str):
        super().__init__()
        self.base_url = base_url
        self.links: list[str] = []
        self.context: dict[str, str] = {}
        self._anchor_url: str | None = None
        self._anchor_text: list[str] = []

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        if tag.lower() != "a":
            return
        href = dict(attrs).get("href")
        if href:
            self._anchor_url = _append_link(self.links, href, self.base_url)
            self._anchor_text = []

    def handle_data(self, data: str) -> None:
        if self._anchor_url is not None:
            self._anchor_text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() != "a" or self._anchor_url is None:
            return
        context = " ".join(" ".join(self._anchor_text).split())
        previous = self.context.get(self._anchor_url, "")
        self.context[self._anchor_url] = " ".join(
            item for item in (previous, context) if item
        )
        self._anchor_url = None
        self._anchor_text = []


def _append_link(links: list[str], raw: str, base_url: str) -> str | None:
    url = urljoin(base_url, raw.strip())
    parsed = urlparse(url)
    if parsed.scheme in ("http", "https") and parsed.hostname:
        url = parsed._replace(fragment="").geturl()
        if url not in links:
            links.append(url)
        return url
    return None


def _link_relevance(text: str, terms: list[str] | None) -> float:
    normalized = text.casefold()
    return float(
        sum(1 for term in terms or [] if term.casefold() in normalized)
    )


def _text_links(text: str, base_url: str) -> list[str]:
    links: list[str] = []
    for url in re.findall(r"https?://[^\s<>\]\[\"']+", text):
        _append_link(links, url.rstrip(".,);"), base_url)
    return links


def _link_relevance_map(
    links: list[str], text: str, terms: list[str] | None
) -> dict[str, float]:
    lines = text.splitlines()
    return {
        link: _link_relevance(
            " ".join(line for line in lines if link in line) or link,
            terms,
        )
        for link in links
    }


def _link_relevance_with_document_context(
    links: list[str], text: str, terms: list[str] | None
) -> dict[str, float]:
    link_context = _link_relevance_map(links, text, terms)
    document_context = _link_relevance(text, terms)
    return {link: max(link_context[link], document_context) for link in links}


def is_rejected_resource(mime_type: str, url: str) -> bool:
    """Reject active content, archives, and executables before processing."""
    mime = mime_type.split(";", 1)[0].strip().lower()
    suffix = Path(urlparse(url).path).suffix.lower()
    return mime in _REJECTED_MIMES or suffix in _REJECTED_SUFFIXES


def is_supported_resource(mime_type: str, url: str) -> bool:
    mime = mime_type.split(";", 1)[0].strip().lower()
    suffix = Path(urlparse(url).path).suffix.lower()
    supported_mime = bool(
        mime
        in {
            "application/pdf",
            "application/xhtml+xml",
            "text/csv",
            "text/html",
            "text/plain",
        }
        or mime.startswith("text/plain")
        or mime in _OFFICE_MIMES
        or mime in _IMAGE_MIMES | _AUDIO_MIMES | _VIDEO_MIMES
    )
    supported_suffix = suffix in (
        _PLAIN_SUFFIXES
        | _IMAGE_SUFFIXES
        | _AUDIO_SUFFIXES
        | _VIDEO_SUFFIXES
        | {".doc", ".docx", ".pdf", ".ppt", ".pptx", ".xls", ".xlsx"}
    )
    return supported_mime or (mime in _GENERIC_MIMES and supported_suffix)


def _number_lines(text: str) -> str:
    return "\n".join(
        f"{number}: {line.strip()}"
        for number, line in enumerate(text.splitlines(), 1)
        if line.strip()
    )


class _ExtractionCancelled(Exception):
    pass


class _ProcessorUnavailable(Exception):
    pass


class _ExtractionLimitExceeded(Exception):
    pass


def _preflight_ooxml(raw: bytes) -> None:
    stream = BytesIO(raw)
    if not zipfile.is_zipfile(stream):
        return
    stream.seek(0)
    with zipfile.ZipFile(stream) as archive:
        members = archive.infolist()
        if len(members) > _MAX_OOXML_MEMBERS:
            raise _ExtractionLimitExceeded("OOXML member limit exceeded")
        if sum(member.file_size for member in members) > (
            _MAX_OOXML_EXPANDED_BYTES
        ):
            raise _ExtractionLimitExceeded(
                "OOXML expanded byte limit exceeded"
            )


def _validate_image_pixels(raw: bytes) -> None:
    if raw.startswith(b"\x89PNG\r\n\x1a\n") and raw[12:16] == b"IHDR":
        width = int.from_bytes(raw[16:20], "big")
        height = int.from_bytes(raw[20:24], "big")
    else:
        image_module = import_module("PIL.Image")
        with image_module.open(BytesIO(raw)) as image:
            width, height = image.size
    if width * height > _MAX_IMAGE_PIXELS:
        raise _ExtractionLimitExceeded("image pixel limit exceeded")


def _terminate_process(process: subprocess.Popen[bytes]) -> None:
    if process.poll() is not None:
        return
    if os.name == "posix" and not _RESOURCE_WORKER_ACTIVE:
        os.killpg(process.pid, signal.SIGTERM)
    else:
        process.terminate()
    try:
        process.wait(timeout=1)
    except subprocess.TimeoutExpired:
        if os.name == "posix" and not _RESOURCE_WORKER_ACTIVE:
            os.killpg(process.pid, signal.SIGKILL)
        else:
            process.kill()
        process.wait()


def _run_process(
    command: list[str], cancellation_event: threading.Event | None = None
) -> subprocess.CompletedProcess[bytes]:
    """Run one owned child process and terminate it when extraction stops."""
    process = subprocess.Popen(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        start_new_session=os.name == "posix" and not _RESOURCE_WORKER_ACTIVE,
    )
    while True:
        if cancellation_event is not None and cancellation_event.is_set():
            _terminate_process(process)
            raise _ExtractionCancelled("extraction cancelled")
        try:
            stdout, stderr = process.communicate(timeout=0.05)
            break
        except subprocess.TimeoutExpired:
            continue
    completed = subprocess.CompletedProcess(
        command, process.returncode, stdout, stderr
    )
    if completed.returncode:
        raise subprocess.CalledProcessError(
            completed.returncode,
            command,
            output=stdout,
            stderr=stderr,
        )
    return completed


def _ocr_image(
    raw: bytes,
    languages: str,
    cancellation_event: threading.Event | None = None,
) -> str:
    _validate_image_pixels(raw)
    executable = shutil.which("tesseract")
    if not executable:
        raise FileNotFoundError("tesseract")
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / "source.img"
        source.write_bytes(raw)
        result = _run_process(
            [executable, str(source), "stdout", "-l", languages],
            cancellation_event,
        )
        return result.stdout.decode("utf-8", errors="replace")


def _extract_pdf(
    raw: bytes,
    languages: str,
    base_url: str,
    cancellation_event: threading.Event | None = None,
) -> tuple[str, list[str], str]:
    import fitz

    document = fitz.open(stream=raw, filetype="pdf")
    try:
        if len(document) > _MAX_PDF_PAGES:
            raise _ExtractionLimitExceeded("PDF page limit exceeded")
        text = "\n".join(str(page.get_text("text")) for page in document)
        links: list[str] = []
        for page in document:
            for link in page.get_links():
                if uri := link.get("uri"):
                    _append_link(links, str(uri), base_url)
        normalized = "\n".join(
            line.strip() for line in text.splitlines() if line.strip()
        )
        if len(re.sub(r"\s", "", normalized)) >= 20:
            return (
                normalized,
                links + _text_links(normalized, base_url),
                "pymupdf",
            )
        pages = []
        raster_pixels = 0
        for page in document:
            page_pixels = round(page.rect.width * 2) * round(
                page.rect.height * 2
            )
            if page_pixels > _MAX_PDF_PAGE_PIXELS:
                raise _ExtractionLimitExceeded(
                    "PDF raster pixel limit exceeded"
                )
            raster_pixels += page_pixels
            if raster_pixels > _MAX_PDF_RASTER_PIXELS:
                raise _ExtractionLimitExceeded(
                    "PDF raster pixel limit exceeded"
                )
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = pixmap.tobytes("png")
            pages.append(
                _ocr_image(image, languages)
                if cancellation_event is None
                else _ocr_image(image, languages, cancellation_event)
            )
        return _number_lines("\n".join(pages)), links, "tesseract"
    finally:
        document.close()


def _extract_docx(raw: bytes, base_url: str) -> tuple[str, list[str]]:
    from docx import Document

    document = Document(BytesIO(raw))
    links: list[str] = []
    for relationship in document.part.rels.values():
        if getattr(relationship, "is_external", False):
            _append_link(links, str(relationship.target_ref), base_url)
    text = extract_docx_text(raw)
    for link in _text_links(text, base_url):
        _append_link(links, link, base_url)
    return text, links


def _extract_xlsx(raw: bytes, base_url: str) -> tuple[str, list[str]]:
    openpyxl = import_module("openpyxl")

    workbook = openpyxl.load_workbook(
        BytesIO(raw), read_only=False, data_only=True
    )
    lines: list[str] = []
    links: list[str] = []
    try:
        if len(workbook.worksheets) > _MAX_SPREADSHEET_SHEETS:
            raise _ExtractionLimitExceeded("spreadsheet sheet limit exceeded")
        declared_cells = sum(
            int(getattr(sheet, "max_row", 0) or 0)
            * int(getattr(sheet, "max_column", 0) or 0)
            for sheet in workbook.worksheets
        )
        if declared_cells > _MAX_SPREADSHEET_CELLS:
            raise _ExtractionLimitExceeded("spreadsheet cell limit exceeded")
        visited_cells = 0
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    visited_cells += 1
                    if visited_cells > _MAX_SPREADSHEET_CELLS:
                        raise _ExtractionLimitExceeded(
                            "spreadsheet cell limit exceeded"
                        )
                    if cell.value is not None:
                        lines.append(
                            f"{sheet.title}!{cell.coordinate}: {cell.value}"
                        )
                    if cell.hyperlink and cell.hyperlink.target:
                        _append_link(links, cell.hyperlink.target, base_url)
    finally:
        workbook.close()
    text = "\n".join(lines)
    return text, links + _text_links(text, base_url)


def _extract_pptx(raw: bytes, base_url: str) -> tuple[str, list[str]]:
    presentation_class = import_module("pptx").Presentation

    presentation = presentation_class(BytesIO(raw))
    if len(presentation.slides) > _MAX_PRESENTATION_SLIDES:
        raise _ExtractionLimitExceeded("presentation slide limit exceeded")
    lines: list[str] = []
    links: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if text := getattr(shape, "text", "").strip():
                lines.append(f"Slide {slide_number}: {text}")
            click = getattr(shape, "click_action", None)
            if click and click.hyperlink.address:
                _append_link(links, click.hyperlink.address, base_url)
    text = "\n".join(lines)
    return text, links + _text_links(text, base_url)


def _convert_legacy_office(
    raw: bytes,
    suffix: str,
    target: str,
    cancellation_event: threading.Event | None = None,
) -> bytes:
    executable = shutil.which("libreoffice")
    if not executable:
        raise FileNotFoundError("libreoffice")
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / f"source.{suffix}"
        source.write_bytes(raw)
        _run_process(
            [
                executable,
                "--headless",
                "--convert-to",
                target,
                "--outdir",
                directory,
                str(source),
            ],
            cancellation_event,
        )
        output = Path(directory) / f"source.{target}"
        if not output.exists():
            raise RuntimeError(f"LibreOffice did not create {target}")
        return output.read_bytes()


def _whisper_worker(audio_path: str, model_name: str, results: Any) -> None:
    """Load and run faster-whisper entirely inside one owned process."""
    try:
        whisper_model_class = import_module("faster_whisper").WhisperModel
        segments, _ = whisper_model_class(model_name).transcribe(audio_path)
        results.put(
            (
                "complete",
                [
                    (
                        float(segment.start),
                        float(segment.end),
                        str(segment.text).strip(),
                    )
                    for segment in segments
                    if str(segment.text).strip()
                ],
            )
        )
    except (ImportError, FileNotFoundError) as error:
        results.put(("unavailable", str(error)))
    except Exception as error:
        results.put(("failed", str(error)))


def _terminate_worker(process: Any) -> None:
    if not process.is_alive():
        process.join()
        return
    process.terminate()
    process.join(timeout=1)
    if process.is_alive():
        process.kill()
        process.join()


def _run_whisper_worker(
    audio: Path,
    model_name: str,
    cancellation_event: threading.Event | None,
) -> list[tuple[float, float, str]]:
    if cancellation_event is not None and cancellation_event.is_set():
        raise _ExtractionCancelled("extraction cancelled")
    context = multiprocessing.get_context("spawn")
    results = context.Queue()
    process = context.Process(
        target=_whisper_worker,
        args=(str(audio), model_name, results),
    )
    process.start()
    try:
        result = None
        while result is None:
            if cancellation_event is not None and cancellation_event.is_set():
                _terminate_worker(process)
                raise _ExtractionCancelled("extraction cancelled")
            try:
                result = results.get(timeout=0.05)
            except queue.Empty:
                if not process.is_alive():
                    try:
                        result = results.get(timeout=1)
                    except queue.Empty as error:
                        raise RuntimeError(
                            "faster-whisper worker exited with code "
                            f"{process.exitcode}"
                        ) from error
        while process.is_alive():
            if cancellation_event is not None and cancellation_event.is_set():
                _terminate_worker(process)
                raise _ExtractionCancelled("extraction cancelled")
            process.join(timeout=0.05)
        status, payload = result
    finally:
        _terminate_worker(process)
        results.close()
        results.cancel_join_thread()
    if status == "unavailable":
        raise _ProcessorUnavailable(str(payload))
    if status == "failed":
        raise RuntimeError(str(payload))
    if status != "complete" or not isinstance(payload, list):
        raise RuntimeError("invalid faster-whisper worker result")
    return [
        (float(start), float(end), str(text)) for start, end, text in payload
    ]


def _transcribe_media(
    raw: bytes,
    suffix: str,
    model_name: str,
    cancellation_event: threading.Event | None = None,
) -> list[tuple[float, float, str]]:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise FileNotFoundError("ffmpeg")
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / f"source{suffix or '.bin'}"
        audio = Path(directory) / "audio.wav"
        source.write_bytes(raw)
        _run_process(
            [
                ffmpeg,
                "-loglevel",
                "error",
                "-i",
                str(source),
                "-vn",
                "-ac",
                "1",
                "-ar",
                "16000",
                str(audio),
            ],
            cancellation_event,
        )
        return _run_whisper_worker(audio, model_name, cancellation_event)


def _timestamp(seconds: float) -> str:
    milliseconds = round(seconds * 1000)
    hours, milliseconds = divmod(milliseconds, 3_600_000)
    minutes, milliseconds = divmod(milliseconds, 60_000)
    whole_seconds, milliseconds = divmod(milliseconds, 1000)
    return f"{hours:02}:{minutes:02}:{whole_seconds:02}.{milliseconds:03}"


def extract_resource(
    raw: bytes,
    mime_type: str,
    url: str,
    *,
    ocr_languages: str = "chi_sim+eng",
    whisper_model: str = "small",
    relevance_terms: list[str] | None = None,
    cancellation_event: threading.Event | None = None,
) -> ExtractionResult:
    """Extract supported content without making optional processors mandatory."""
    mime = mime_type.split(";", 1)[0].strip().lower()
    suffix = Path(urlparse(url).path).suffix.lower()
    if is_rejected_resource(mime, url):
        return ExtractionResult(
            status="skipped", error="rejected resource type"
        )
    if not is_supported_resource(mime, url):
        return ExtractionResult(
            status="skipped", error="unsupported resource type"
        )
    try:
        if mime in {"text/html", "application/xhtml+xml"} or suffix in {
            ".htm",
            ".html",
        }:
            decoded = decode_body(raw, mime_type)
            parser = _LinkParser(url)
            parser.feed(decoded)
            extracted = extract_html(decoded)
            return ExtractionResult(
                status="complete",
                text=extracted["text"],
                links=parser.links,
                link_relevance={
                    link: _link_relevance(
                        f"{link} {parser.context.get(link, '')}",
                        relevance_terms,
                    )
                    for link in parser.links
                },
                processor="html",
            )
        if mime == "application/pdf" or suffix == ".pdf":
            text, links, processor = _extract_pdf(
                raw, ocr_languages, url, cancellation_event
            )
            return ExtractionResult(
                status="complete",
                text=text,
                links=links,
                link_relevance=_link_relevance_with_document_context(
                    links, text, relevance_terms
                ),
                processor=processor,
            )
        office_type = _OFFICE_MIMES.get(mime) or suffix.lstrip(".")
        if office_type == "doc":
            raw = _convert_legacy_office(
                raw, "doc", "docx", cancellation_event
            )
            office_type = "docx"
        elif office_type == "xls":
            raw = _convert_legacy_office(
                raw, "xls", "xlsx", cancellation_event
            )
            office_type = "xlsx"
        elif office_type == "ppt":
            raw = _convert_legacy_office(
                raw, "ppt", "pptx", cancellation_event
            )
            office_type = "pptx"
        if office_type in {"docx", "xlsx", "pptx"}:
            _preflight_ooxml(raw)
        if office_type == "docx":
            text, links = _extract_docx(raw, url)
            return ExtractionResult(
                status="complete",
                text=text,
                links=links,
                link_relevance=_link_relevance_with_document_context(
                    links, text, relevance_terms
                ),
                processor="python-docx",
            )
        if office_type == "xlsx":
            text, links = _extract_xlsx(raw, url)
            return ExtractionResult(
                status="complete",
                text=text,
                links=links,
                link_relevance=_link_relevance_with_document_context(
                    links, text, relevance_terms
                ),
                processor="openpyxl",
            )
        if office_type == "pptx":
            text, links = _extract_pptx(raw, url)
            return ExtractionResult(
                status="complete",
                text=text,
                links=links,
                link_relevance=_link_relevance_with_document_context(
                    links, text, relevance_terms
                ),
                processor="python-pptx",
            )
        if mime in _IMAGE_MIMES or (
            mime in _GENERIC_MIMES and suffix in _IMAGE_SUFFIXES
        ):
            return ExtractionResult(
                status="complete",
                text=_number_lines(
                    _ocr_image(raw, ocr_languages)
                    if cancellation_event is None
                    else _ocr_image(raw, ocr_languages, cancellation_event)
                ),
                processor="tesseract",
            )
        if mime in _AUDIO_MIMES | _VIDEO_MIMES or (
            mime in _GENERIC_MIMES
            and suffix in _AUDIO_SUFFIXES | _VIDEO_SUFFIXES
        ):
            segments = (
                _transcribe_media(raw, suffix, whisper_model)
                if cancellation_event is None
                else _transcribe_media(
                    raw, suffix, whisper_model, cancellation_event
                )
            )
            text = "\n".join(
                f"[{_timestamp(start)} --> {_timestamp(end)}] {segment}"
                for start, end, segment in segments
            )
            return ExtractionResult(
                status="complete", text=text, processor="faster-whisper"
            )
        text = decode_body(raw, mime_type).replace("\r\n", "\n").strip()
        links = _text_links(text, url)
        return ExtractionResult(
            status="complete",
            text=text,
            links=links,
            link_relevance=_link_relevance_map(links, text, relevance_terms),
            processor="text",
        )
    except (
        ImportError,
        FileNotFoundError,
        _ProcessorUnavailable,
        subprocess.SubprocessError,
    ) as error:
        return ExtractionResult(status="unavailable", error=str(error))
    except Exception as error:
        return ExtractionResult(status="failed", error=str(error))


def _resource_worker(
    raw: bytes,
    mime_type: str,
    url: str,
    ocr_languages: str,
    whisper_model: str,
    relevance_terms: list[str] | None,
    results: Any,
) -> None:
    global _RESOURCE_WORKER_ACTIVE
    if os.name == "posix":
        os.setsid()
    _RESOURCE_WORKER_ACTIVE = True
    try:
        result = extract_resource(
            raw,
            mime_type,
            url,
            ocr_languages=ocr_languages,
            whisper_model=whisper_model,
            relevance_terms=relevance_terms,
        )
        results.put(result.model_dump())
    except Exception as error:
        results.put(
            ExtractionResult(status="failed", error=str(error)).model_dump()
        )


def _terminate_extraction_worker(process: Any) -> None:
    if not process.is_alive():
        process.join()
        return
    if os.name == "posix" and process.pid is not None:
        try:
            os.killpg(process.pid, signal.SIGTERM)
        except ProcessLookupError:
            process.terminate()
    else:
        process.terminate()
    process.join(timeout=1)
    if process.is_alive():
        if os.name == "posix" and process.pid is not None:
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                process.kill()
        else:
            process.kill()
        process.join()


def extract_resource_process(
    raw: bytes,
    mime_type: str,
    url: str,
    *,
    ocr_languages: str = "chi_sim+eng",
    whisper_model: str = "small",
    relevance_terms: list[str] | None = None,
    cancellation_event: threading.Event | None = None,
) -> ExtractionResult:
    """Run document parsers in a worker that can be killed or timed out."""
    mime = mime_type.split(";", 1)[0].strip().lower()
    suffix = Path(urlparse(url).path).suffix.lower()
    if not (
        mime == "application/pdf"
        or mime in _OFFICE_MIMES
        or mime in _IMAGE_MIMES
        or suffix
        in _IMAGE_SUFFIXES
        | {".doc", ".docx", ".pdf", ".ppt", ".pptx", ".xls", ".xlsx"}
    ):
        return extract_resource(
            raw,
            mime_type,
            url,
            ocr_languages=ocr_languages,
            whisper_model=whisper_model,
            relevance_terms=relevance_terms,
            cancellation_event=cancellation_event,
        )
    context = multiprocessing.get_context("spawn")
    results = context.Queue()
    process = context.Process(
        target=_resource_worker,
        args=(
            raw,
            mime_type,
            url,
            ocr_languages,
            whisper_model,
            relevance_terms,
            results,
        ),
    )
    process.start()
    deadline = time.monotonic() + _EXTRACTION_TIMEOUT_SECONDS
    payload: object | None = None
    try:
        while payload is None:
            if cancellation_event is not None and cancellation_event.is_set():
                raise _ExtractionCancelled("extraction cancelled")
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                return ExtractionResult(
                    status="failed", error="extraction timed out"
                )
            try:
                payload = results.get(timeout=min(0.05, remaining))
            except queue.Empty:
                if not process.is_alive():
                    try:
                        payload = results.get(timeout=0.1)
                    except queue.Empty:
                        return ExtractionResult(
                            status="failed",
                            error=(
                                "extraction worker exited with code "
                                f"{process.exitcode}"
                            ),
                        )
        while process.is_alive():
            if cancellation_event is not None and cancellation_event.is_set():
                raise _ExtractionCancelled("extraction cancelled")
            if time.monotonic() >= deadline:
                return ExtractionResult(
                    status="failed", error="extraction timed out"
                )
            process.join(timeout=0.05)
    finally:
        _terminate_extraction_worker(process)
        results.close()
        results.cancel_join_thread()
    return ExtractionResult.model_validate(payload)
