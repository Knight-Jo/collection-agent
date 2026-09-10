"""Report export renderer and reference extraction tests."""

from __future__ import annotations

import io
from datetime import UTC, datetime

import pymupdf
import pytest
from docx import Document

from intel_agent.contracts.errors import DomainError
from intel_agent.contracts.research import ReportSection, ResearchReport
from intel_agent.reporting import (
    ExportMeta,
    ascii_slug,
    export_filename,
    normalize_format,
    references_for_report,
    references_from_evidence,
    render_docx,
    render_markdown,
    render_pdf,
    render_report,
)

META = ExportMeta(
    subject="测试主题", generated_at=datetime(2026, 9, 8, 8, 0, tzinfo=UTC)
)

REPORT = ResearchReport(
    title="Acme 调研报告",
    sections=[
        ReportSection(
            heading="核心发现",
            body="**关键**结论如下：\n- 要点一\n- 要点二\n1. 编号项",
        )
    ],
    conclusions=["总体向好"],
    limitations=["样本有限"],
    citation_ids=["C1"],
)

EVIDENCE = {
    "claims": [
        {
            "claim": "a",
            "citation_id": "C1",
            "source_title": "来源甲",
            "source_url": "https://example.com/a",
        },
        {
            "claim": "b",
            "citation_id": "C1",
            "source_title": "来源甲",
            "source_url": "https://example.com/a",  # duplicate source
        },
        {
            "claim": "c",
            "citation_id": "C2",
            "source_title": "来源乙",
            "source_url": "https://example.com/b",
        },
        {
            "claim": "d",
            "citation_id": "C3",
            "source_title": "来源丙",
            "source_url": "https://example.com/c",
        },
    ]
}

REPORT_CJK_TITLE = REPORT.model_copy(update={"title": "纯中文标题"})


def test_references_are_citation_keyed():
    # the report cites C2 first; cited keys lead, uncited C3 follows
    report = REPORT.model_copy(update={"citation_ids": ["C2", "C1"]})
    refs = references_for_report(report, EVIDENCE)
    assert [(r.key, r.title) for r in refs] == [
        ("C2", "来源乙"),
        ("C1", "来源甲"),
        ("C3", "来源丙"),  # present in evidence, not cited: still listed
    ]


def test_references_unresolvable_key_gets_placeholder():
    report = REPORT.model_copy(update={"citation_ids": ["C1", "C9"]})
    refs = references_for_report(report, EVIDENCE)
    dangling = refs[1]
    assert dangling.key == "C9"
    assert dangling.title == "（来源信息缺失）"
    assert dangling.url == ""


def test_references_without_citation_ids_cover_all_sources():
    report = REPORT.model_copy(update={"citation_ids": []})
    refs = references_for_report(report, EVIDENCE)
    assert {r.key for r in refs} == {"C1", "C2", "C3"}


def test_references_dedup_and_order():
    refs = references_from_evidence(EVIDENCE)
    assert [(r.key, r.title, r.url) for r in refs] == [
        ("C1", "来源甲", "https://example.com/a"),
        ("C2", "来源乙", "https://example.com/b"),
        ("C3", "来源丙", "https://example.com/c"),
    ]
    assert references_from_evidence(None) == []
    assert references_from_evidence({"claims": []}) == []


def test_markdown_contains_header_sections_and_references():
    text = render_markdown(REPORT, references_from_evidence(EVIDENCE), META)
    decoded = text.decode("utf-8")
    assert decoded.startswith("# Acme 调研报告")
    assert "主题：测试主题" in decoded
    assert "## 核心发现" in decoded
    assert "- 要点一" in decoded
    assert "## 结论" in decoded
    assert "## 参考文献" in decoded
    assert "[C1] [来源甲](https://example.com/a)" in decoded
    assert decoded.count("来源甲") == 1  # deduplicated


def _style(paragraph) -> str:
    style = paragraph.style
    return style.name if style is not None else ""


def test_docx_structure_round_trip():
    payload = render_docx(REPORT, references_from_evidence(EVIDENCE), META)
    doc = Document(io.BytesIO(payload))
    headings = [
        p.text
        for p in doc.paragraphs
        if _style(p).startswith(("Heading", "Title"))
    ]
    assert "Acme 调研报告" in headings
    assert "核心发现" in headings
    assert "参考文献" in headings
    bullets = [p.text for p in doc.paragraphs if _style(p) == "List Bullet"]
    assert "要点一" in bullets and "要点二" in bullets
    reference_par = [
        p.text for p in doc.paragraphs if p.text.startswith("[C1]")
    ]
    assert any("来源甲" in text for text in reference_par)
    # bold run survives the markdown conversion
    bold_runs = [
        run.text
        for p in doc.paragraphs
        for run in p.runs
        if run.bold and run.text.strip()
    ]
    assert any("关键" in text for text in bold_runs)


def test_pdf_renders_cjk_text():
    payload = render_pdf(REPORT, references_from_evidence(EVIDENCE), META)
    assert payload.startswith(b"%PDF")
    doc = pymupdf.open(stream=payload, filetype="pdf")
    assert doc.page_count >= 1
    text = "\n".join(str(page.get_text()) for page in doc)
    for probe in (
        "Acme 调研报告",
        "核心发现",
        "要点一",
        "总体向好",
        "参考文献",
        "example.com/a",
    ):
        assert probe in text


def test_normalize_format_aliases_and_errors():
    assert normalize_format("md") == "markdown"
    assert normalize_format("PDF ") == "pdf"
    assert normalize_format("word") == "docx"
    with pytest.raises(ValueError, match="xlsx"):
        normalize_format("xlsx")


def test_render_report_dispatches_formats():
    for fmt in ("markdown", "pdf", "docx"):
        payload = render_report(REPORT, EVIDENCE, fmt, subject="测试主题")
        assert isinstance(payload, bytes) and payload
    assert render_report(REPORT, EVIDENCE, "md").startswith(b"# ")


def test_filename_is_ascii_and_dated():
    # CJK-only titles fall back to a safe ASCII stem
    name = export_filename(REPORT_CJK_TITLE, "pdf", META)
    assert name == "report-20260908.pdf"
    ascii_report = REPORT.model_copy(update={"title": "Acme Report"})
    assert (
        export_filename(ascii_report, "docx", META)
        == "acme-report-20260908.docx"
    )
    assert ascii_slug("Hello World!!") == "hello-world"


# --- conversation service export path ---------------------------------------


class _StubStore:
    """Minimal MaterialStore stand-in for the export path."""

    def __init__(self, result: dict | None) -> None:
        self._result = result

    def latest_task_id(self, conversation_id: str) -> str | None:
        return "task-1" if self._result else None

    def get_research_result(self, task_id: str) -> dict | None:
        return self._result

    def get_conversation(self, conversation_id: str) -> dict:
        return {"id": conversation_id, "title": "会话标题"}


def _export_service(result: dict | None):
    from intel_agent.conversation import ConversationService

    stub = _StubStore(result)
    return ConversationService(
        store=stub,  # type: ignore[arg-type]
        orchestrator=None,  # type: ignore[arg-type]
        event_bus=None,  # type: ignore[arg-type]
        roles={},
        registry=None,  # type: ignore[arg-type]
        settings=None,  # type: ignore[arg-type]
        task_store=stub,  # type: ignore[arg-type]
    )


def _stored_result() -> dict:
    return {
        "report": REPORT.model_dump(mode="json"),
        "coverage": None,
        "evidence": EVIDENCE,
    }


def test_export_report_returns_payload_media_type_filename():
    service = _export_service(_stored_result())
    for fmt, prefix, media in (
        ("markdown", b"# ", "text/markdown"),
        ("pdf", b"%PDF", "application/pdf"),
        ("docx", b"PK", "wordprocessingml"),
    ):
        payload, media_type, filename = service.export_report("c1", fmt)
        assert payload.startswith(prefix)
        assert media in media_type or media_type.startswith(
            media.split(";")[0]
        )
        assert filename.endswith(
            ("md", "pdf", "docx")[("markdown", "pdf", "docx").index(fmt)]
        )
    # alias works too
    payload, _media, _name = service.export_report("c1", "word")
    assert payload.startswith(b"PK")


def test_export_report_errors():
    service = _export_service(None)  # no result row
    with pytest.raises(DomainError) as raised:
        service.export_report("c1", "pdf")
    assert raised.value.code == "NOT_FOUND"

    empty = _export_service(
        {"report": None, "coverage": None, "evidence": None}
    )
    with pytest.raises(DomainError) as raised:
        empty.export_report("c1", "pdf")
    assert raised.value.code == "NOT_FOUND"

    with pytest.raises(DomainError) as raised:
        _export_service(_stored_result()).export_report("c1", "rtf")
    assert raised.value.code == "INVALID_REQUEST"
